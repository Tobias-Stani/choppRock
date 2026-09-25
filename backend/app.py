import hashlib
import hmac
import os
import sqlite3
from datetime import datetime
from pathlib import Path

from fastapi import Cookie, Depends, FastAPI, File, HTTPException, Response, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel

CLIENT_CODE = os.environ.get("CLIENT_CODE", "rock")  # código inicial; después se cambia desde el admin
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin")
DB = Path(os.environ.get("DB_PATH", "catalogo.db"))
PAGE_SIZE = 50
MONTH = 60 * 60 * 24 * 30

# encabezado del Excel (en minúsculas) -> columna en la base; se matchea por prefijo
COLUMNS = {"artist": "artist", "title": "title", "label": "label",
           "med": "media", "desc": "description", "genre": "genre"}

app = FastAPI()


class Secret(BaseModel):
    value: str


def db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con


def get_setting(key, default=None):
    con = db()
    con.execute("CREATE TABLE IF NOT EXISTS settings (key PRIMARY KEY, value)")
    row = con.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    con.close()
    return row[0] if row else default


def set_settings(**values):
    con = db()
    with con:
        con.execute("CREATE TABLE IF NOT EXISTS settings (key PRIMARY KEY, value)")
        con.executemany("INSERT OR REPLACE INTO settings VALUES (?, ?)", values.items())
    con.close()


def client_code():
    return get_setting("client_code", CLIENT_CODE)


def sign(value):
    # cookie = firma, no el secreto; cambiar el código invalida las sesiones de clientes
    return hmac.new(ADMIN_PASSWORD.encode(), value.encode(), hashlib.sha256).hexdigest()


def same(a, b):
    return bool(a) and hmac.compare_digest(a, b)


def require_admin(admin: str | None = Cookie(None)):
    if not same(admin, sign("admin")):
        raise HTTPException(401, "Iniciá sesión como administrador.")


def require_client(session: str | None = Cookie(None), admin: str | None = Cookie(None)):
    if not (same(session, sign("client:" + client_code())) or same(admin, sign("admin"))):
        raise HTTPException(401, "Ingresá el código de acceso.")


def has_catalog():
    con = db()
    ok = con.execute("SELECT 1 FROM sqlite_master WHERE name = 'discos'").fetchone()
    con.close()
    return bool(ok)


def load_excel(fileobj):
    try:
        ws = load_workbook(fileobj, read_only=True).active
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo. Tiene que ser un Excel .xlsx.")
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows, [])]
    idx = {}
    for i, h in enumerate(header):
        for prefix, col in COLUMNS.items():
            if h.startswith(prefix) and col not in idx:
                idx[col] = i
    if "artist" not in idx or "title" not in idx:
        raise HTTPException(400, "Falta la columna Artist o Title en la primera fila del Excel.")

    out = []
    # id = número de fila del Excel, así la disquería lo encuentra directo
    for n, row in enumerate(rows, start=2):
        rec = {c: str(row[i]).strip() if i < len(row) and row[i] is not None else ""
               for c, i in idx.items()}
        if rec["artist"] or rec["title"]:
            out.append((n, *(rec.get(c, "") for c in COLUMNS.values())))
    return out


def save(rows):
    con = db()
    with con:
        con.execute("DROP TABLE IF EXISTS discos")
        con.execute("CREATE TABLE discos (id INTEGER PRIMARY KEY, artist, title, label, media, description, genre)")
        con.executemany("INSERT INTO discos VALUES (?,?,?,?,?,?,?)", rows)
    con.close()


# --- clientes ---

@app.post("/api/login")
def login(body: Secret, response: Response):
    if not same(body.value.strip(), client_code()):
        raise HTTPException(401, "Ese código no es correcto. Pedíselo a la disquería.")
    response.set_cookie("session", sign("client:" + client_code()), httponly=True, samesite="lax", max_age=MONTH)
    return {"ok": True}


@app.post("/api/logout")
def logout(response: Response):
    response.delete_cookie("session")
    return {"ok": True}


@app.get("/api/filtros", dependencies=[Depends(require_client)])
def filtros():
    if not has_catalog():
        return {"media": [], "genre": []}
    con = db()
    res = {c: [r[0] for r in con.execute(f"SELECT DISTINCT {c} FROM discos WHERE {c} != '' ORDER BY {c}")]
           for c in ("media", "genre")}
    con.close()
    return res


@app.get("/api/discos", dependencies=[Depends(require_client)])
def discos(q: str = "", media: str = "", genre: str = "", page: int = 1):
    if not has_catalog():
        return {"total": 0, "items": [], "page_size": PAGE_SIZE}
    where, params = [], []
    # cada palabra tiene que aparecer en artista, título o sello
    # ponytail: LIKE full scan, ~60k filas va sobrado; FTS5 si el catálogo crece mucho
    for word in q.split():
        where.append("(artist || ' ' || title || ' ' || label) LIKE ?")
        params.append(f"%{word}%")
    if media:
        where.append("media = ?"); params.append(media)
    if genre:
        where.append("genre = ?"); params.append(genre)
    sql_where = f"WHERE {' AND '.join(where)}" if where else ""
    con = db()
    total = con.execute(f"SELECT COUNT(*) FROM discos {sql_where}", params).fetchone()[0]
    items = [dict(r) for r in con.execute(
        f"SELECT * FROM discos {sql_where} ORDER BY id LIMIT ? OFFSET ?",
        [*params, PAGE_SIZE, (max(page, 1) - 1) * PAGE_SIZE])]
    con.close()
    return {"total": total, "items": items, "page_size": PAGE_SIZE}


# --- admin ---

@app.post("/api/admin/login")
def admin_login(body: Secret, response: Response):
    if not same(body.value, ADMIN_PASSWORD):
        raise HTTPException(401, "Contraseña incorrecta.")
    response.set_cookie("admin", sign("admin"), httponly=True, samesite="strict", max_age=MONTH)
    return {"ok": True}


@app.post("/api/admin/logout")
def admin_logout(response: Response):
    response.delete_cookie("admin")
    return {"ok": True}


@app.get("/api/admin/status", dependencies=[Depends(require_admin)])
def admin_status():
    res = {"client_code": client_code(), "filename": get_setting("filename"),
           "uploaded_at": get_setting("uploaded_at"), "total": 0, "media": [], "genres": 0}
    if has_catalog():
        con = db()
        res["total"] = con.execute("SELECT COUNT(*) FROM discos").fetchone()[0]
        res["media"] = [dict(r) for r in con.execute(
            "SELECT media, COUNT(*) n FROM discos GROUP BY media ORDER BY n DESC")]
        res["genres"] = con.execute("SELECT COUNT(DISTINCT genre) FROM discos").fetchone()[0]
        con.close()
    return res


@app.post("/api/admin/upload", dependencies=[Depends(require_admin)])
def upload(file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "El archivo tiene que ser .xlsx.")
    rows = load_excel(file.file)
    if not rows:
        raise HTTPException(400, "El Excel no tiene discos. Revisá que sea la lista correcta.")
    save(rows)
    set_settings(filename=file.filename, uploaded_at=datetime.now().isoformat(timespec="minutes"))
    return {"ok": True, "total": len(rows)}


@app.put("/api/admin/code", dependencies=[Depends(require_admin)])
def change_code(body: Secret):
    code = body.value.strip()
    if not 4 <= len(code) <= 64:
        raise HTTPException(400, "El código tiene que tener entre 4 y 64 caracteres.")
    set_settings(client_code=code)
    return {"ok": True, "client_code": code}


# producción (Railway): un solo contenedor sirve también el front. En local lo hace nginx.
if STATIC_DIR := os.environ.get("STATIC_DIR"):
    @app.get("/admin")
    def admin_page():
        return FileResponse(Path(STATIC_DIR) / "admin.html")

    app.mount("/", StaticFiles(directory=STATIC_DIR, html=True))


if __name__ == "__main__":
    # self-check del parser: python app.py
    import io
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["Artist", "Title", "Label", "", "Media", "Description", "Genre"])
    ws.append(["MURPHY, Peter", "Deep", "Beggars", "", "Vinyl", "LP", "Indie"])
    ws.append([None, None, None, None, None, None, None])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    rows = load_excel(buf)
    assert rows == [(2, "MURPHY, Peter", "Deep", "Beggars", "Vinyl", "LP", "Indie")], rows
    print("ok")
